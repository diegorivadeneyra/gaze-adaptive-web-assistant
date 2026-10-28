"""Lee los Excel respondidos (data/respuestas/P01.xlsx…; celdas según la hoja _mapa)
y calcula los puntajes en data/sessions/resumen/:
  comprension_textos.csv  texto por página + columnas para calificar a mano
                          (ideas_correctas, ideas_incorrectas, notas; se conservan).
                          El Excel pregunta por posición ("Página web 1 = la primera que
                          leíste"); la página sale del orden del participante: primero
                          evaluation/asignacion_participantes.csv, si no, su sesión en data/sessions
  nasa_tlx.csv            una fila por participante y método (Método 1 = "ayuda", Método 2 =
                          "a_pedido"): 6 dimensiones 0–100 (Desempeño: 0 = Muy bueno); peso = veces
                          elegida en los 15 pares. bruto = promedio; ponderado = suma(nivel × peso) / 15
  post_experimento.csv    Anexo 4 adaptado (versión final editada el 2026-10-28: 4 órdenes,
                          7 calificaciones; en "aprender" solo se puede elegir 1 o 5)
  post_metodos.csv        una fila por participante y método (incluye el NASA de ese método)
El SUS se quitó el 2026-10-27. Las celdas de cada respuesta están en la hoja oculta
_mapa del Excel (evaluation/cuestionarios_participante.xlsx es la versión final).
  avisos.txt              respuestas vacías o incoherentes

Uso: python evaluation/puntuar_cuestionarios.py
"""
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
IN_DIR = ROOT / "data" / "respuestas"
OUT_DIR = ROOT / "data" / "sessions" / "resumen"
ASSIGNMENT_CSV = ROOT / "evaluation" / "asignacion_participantes.csv"
SESSIONS_DIR = ROOT / "data" / "sessions"

# Hojas del Excel (ver crear_cuestionarios_excel.py).
COMPRENSION_SHEET = "Comprensión"
NASA_SHEETS = {"m1": "NASA-TLX Método 1", "m2": "NASA-TLX Método 2"}
PAGES = ["facil", "intermedio", "avanzado"]
MANUAL = ["ideas_correctas", "ideas_incorrectas", "notas"]   # columnas que se llenan a mano
NASA_KEYS = ["mental", "fisica", "temporal", "desempeno", "esfuerzo", "frustracion"]
NASA_NAME_TO_KEY = {"Exigencia mental": "mental", "Exigencia física": "fisica", "Exigencia temporal": "temporal",
                    "Desempeño": "desempeno", "Esfuerzo": "esfuerzo", "Frustración": "frustracion"}
METHODS = {"m1": "ayuda", "m2": "a_pedido"}          # método del Anexo 4 -> condición del sistema
# Mismo orden que en la hoja Post-experimento (claves de _mapa).
RANKINGS = ["rapidez", "comodo", "control", "elegiria"]   # ayuda más rápida, más cómodo, más control, elegiría
LIKERT = ["completar", "comodidad", "aprender", "rapidez", "precision", "confianza", "futuro"]
COMMENTS = ["positivos", "dificultades", "sugerencias"]


def reading_orders():
    """Orden de páginas de cada participante: el sorteo previo y, para códigos
    fuera de él, el session_start de su sesión (la más reciente)."""
    orders = {}
    for path in sorted(SESSIONS_DIR.glob("*_events.jsonl")):
        with path.open(encoding="utf-8") as f:
            first = f.readline()
        try:
            start = json.loads(first)
        except ValueError:
            continue
        if start.get("type") == "session_start" and start.get("participant") and start.get("order"):
            orders[str(start["participant"]).strip().upper()] = start["order"]
    if ASSIGNMENT_CSV.exists():
        with ASSIGNMENT_CSV.open(encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                orders[r["participante"].strip().upper()] = [r["posicion_1"], r["posicion_2"], r["posicion_3"]]
    return orders


def number(value, low, high):
    """Número de la respuesta; acepta las opciones con significado ("4 = De acuerdo")."""
    match = re.match(r"\s*(\d+(?:\.\d+)?)", str(value)) if value is not None else None
    if not match:
        return None
    v = float(match.group(1))
    return v if low <= v <= high else None


class Answers:
    """Lee las respuestas de un Excel usando su hoja oculta _mapa."""

    def __init__(self, wb):
        self.wb = wb
        self.cells = {(h, k): c for h, k, c in wb["_mapa"].iter_rows(min_row=2, values_only=True)}

    def get(self, sheet, key):
        ref = self.cells.get((sheet, key))
        if ref is None or sheet not in self.wb.sheetnames:
            return None
        value = self.wb[sheet][ref].value
        return None if value is None or str(value).strip() == "" else value


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")   # la consola de Windows no imprime "⚠"
    # Calificaciones a mano ya hechas: se conservan al regenerar el archivo.
    graded = {}
    texts_csv = OUT_DIR / "comprension_textos.csv"
    if texts_csv.exists():
        with texts_csv.open(encoding="utf-8-sig", newline="") as f:
            graded = {(r["participante"], r["pagina"]): r for r in csv.DictReader(f)}
    files = sorted(f for f in IN_DIR.glob("*.xlsx") if not f.name.startswith("~$"))   # ~$ = abierto en Excel
    orders = reading_orders()
    if not files:
        sys.exit(f"No hay Excel en {IN_DIR.relative_to(ROOT)} (guarda cada uno como P01.xlsx, P02.xlsx…)")

    pages, nasa_rows, post_rows, method_rows, warnings = [], [], [], [], []
    for path in files:
        participant = path.stem.strip().upper()
        try:
            wb = load_workbook(path, data_only=True)
        except Exception as error:
            warnings.append(f"{path.name}: no se pudo abrir ({error})")
            continue
        if "_mapa" not in wb.sheetnames:
            warnings.append(f"{path.name}: es de una versión anterior del cuestionario (sin hoja _mapa); no se leyó")
            continue
        ans = Answers(wb)
        warn = lambda msg: warnings.append(f"{participant}, {msg}")

        # ── Comprensión: lo que escribió de cada página ─────────────
        order = orders.get(participant)
        by_position = (COMPRENSION_SHEET, "resumen_pos1") in ans.cells
        if by_position and not order:
            warn("comprensión: sin orden de páginas (no está en el sorteo ni hay sesión); "
                 "se guarda por posición (pos1, pos2, pos3)")
        for i, page in enumerate(PAGES):
            if not by_position:                       # Excel anterior: una clave por página
                key = f"resumen_{page}"
            elif order:                               # "Página web n" -> la n-ésima que leyó
                key = f"resumen_pos{order.index(page) + 1}"
            else:
                key, page = f"resumen_pos{i + 1}", f"pos{i + 1}"
            raw = ans.get(COMPRENSION_SHEET, key)
            text = "" if raw is None else str(raw).strip()
            if not text:
                warn(f"comprensión, página {page}: recuadro vacío")
            old = graded.get((participant, page), {})
            pages.append({"participante": participant, "pagina": page, "palabras": len(re.findall(r"\w+", text)),
                          "texto": text, **{c: old.get(c, "") for c in MANUAL}})

        # ── NASA-TLX: uno por método ────────────────────────────────
        nasa_by_method = {}
        for m, sheet in NASA_SHEETS.items():
            nasa = {"participante": participant, "metodo": m, "condicion": METHODS[m]}
            levels = {}
            for dim in NASA_KEYS:
                v = number(ans.get(sheet, f"nasa_{dim}"), 0, 100)
                if v is None:
                    warn(f"{sheet} '{dim}': sin responder o fuera de 0–100")
                levels[dim] = v
                nasa[f"nasa_{dim}"] = "" if v is None else v
            weights = dict.fromkeys(NASA_KEYS, 0)
            pairs_ok = 0
            for n in range(1, 16):
                choice = ans.get(sheet, f"par_{n:02d}")
                dim = NASA_NAME_TO_KEY.get(str(choice).strip()) if choice is not None else None
                if dim is None:
                    warn(f"{sheet} comparación {n}: sin responder")
                    continue
                weights[dim] += 1
                pairs_ok += 1
            for dim in NASA_KEYS:
                nasa[f"peso_{dim}"] = weights[dim] if pairs_ok == 15 else ""
            complete = None not in levels.values()
            nasa["nasa_bruto"] = round(sum(levels.values()) / 6, 1) if complete else ""
            nasa["nasa_ponderado"] = (round(sum(levels[d] * weights[d] for d in NASA_KEYS) / 15, 1)
                                      if complete and pairs_ok == 15 else "")
            nasa_rows.append(nasa)
            nasa_by_method[m] = nasa

        # ── Post-experimento (Anexo 4) ──────────────────────────────
        post = {"participante": participant}
        per_method = {m: {"participante": participant, "metodo": m, "condicion": c,
                          "nasa_bruto": nasa_by_method[m]["nasa_bruto"],
                          "nasa_ponderado": nasa_by_method[m]["nasa_ponderado"]} for m, c in METHODS.items()}
        for item in RANKINGS:
            ranks = {}
            for m in METHODS:
                v = number(ans.get("Post-experimento", f"orden_{item}_{m}"), 1, 2)
                ranks[m] = None if v is None else int(v)
                post[f"orden_{item}_{m}"] = "" if v is None else int(v)
                per_method[m][f"orden_{item}"] = "" if v is None else int(v)
            if None in ranks.values():
                warn(f"post-experimento, orden '{item}': sin responder")
            elif ranks["m1"] == ranks["m2"]:
                warn(f"post-experimento, orden '{item}': los dos métodos tienen el mismo lugar ({ranks['m1']})")
        for item in LIKERT:
            for m in METHODS:
                v = number(ans.get("Post-experimento", f"califica_{item}_{m}"), 1, 5)
                if v is None:
                    warn(f"post-experimento, '{item}' ({m}): sin responder o fuera de 1–5")
                post[f"califica_{item}_{m}"] = "" if v is None else int(v)
                per_method[m][f"califica_{item}"] = "" if v is None else int(v)
        for item in COMMENTS:
            text = ans.get("Post-experimento", f"comentario_{item}")
            post[f"comentario_{item}"] = "" if text is None else str(text).strip()
        post_rows.append(post)
        method_rows.extend(per_method.values())

    codes = [f.stem.strip().upper() for f in files]
    for code in sorted({c for c in codes if codes.count(c) > 1}):
        warnings.append(f"Código repetido: {code} aparece en más de un archivo")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    outputs = {"comprension_textos.csv": pages, "nasa_tlx.csv": nasa_rows,
               "post_experimento.csv": post_rows, "post_metodos.csv": method_rows}
    for name, rows in outputs.items():
        write(OUT_DIR / name, rows)
    (OUT_DIR / "avisos.txt").write_text("\n".join(warnings) or "Sin avisos.", encoding="utf-8")

    print(f"Leídos {len(files)} archivos: {', '.join(codes)}")
    for row in pages:
        print(f"  {row['participante']:6} {row['pagina']:10} comprensión: {row['palabras']} palabras"
              + (f", calificada ({row['ideas_correctas']} ideas correctas)" if row["ideas_correctas"] != "" else ""))
    dash = lambda v: v if v != "" else "—"
    for nasa in nasa_rows:
        print(f"  {nasa['participante']:6} NASA {nasa['metodo']} ({nasa['condicion']}): bruto {dash(nasa['nasa_bruto'])}, "
              f"ponderado {dash(nasa['nasa_ponderado'])}")
    print(f"\n{len(warnings)} aviso(s)" + (":" if warnings else "."))
    for w in warnings:
        print("  ⚠", w)
    print(f"\nGuardado en {OUT_DIR.relative_to(ROOT)}: {', '.join(outputs)}, avisos.txt")


def write(path, rows):
    if not rows:
        return
    with path.open("w", newline="", encoding="utf-8-sig") as f:   # utf-8-sig: Excel abre bien los acentos
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


if __name__ == "__main__":
    main()
